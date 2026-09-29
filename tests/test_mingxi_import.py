# -*- coding: utf-8 -*-
"""明细导入模块单元测试。"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import config
from processors.mingxi_import import build_mingxi, validate, _lookup_catalog
from processors.tax_rate_resolver import _norm_rate, resolve_tax_rate, save_default_rate, load_default_rate
from processors.optional_col_resolver import resolve_optional_cols, load_prefs, save_prefs
from utils.excel_reader import match_column, read_input, normalize_records


def _make_input(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["序号", "品名", "单位", "数量", "单价", "金额"])
    ws.append([1, "320偏心环", "件", 30, 54, 1620])
    ws.append([2, "255原点盖", "件", 345, 28, 9660])
    wb.save(path)


class TestRateNorm:
    def test_norm(self):
        assert _norm_rate("0.13") == "0.13"
        assert _norm_rate("13") == "0.13"
        assert _norm_rate("13%") == "0.13"
        assert _norm_rate(0.13) == "0.13"
        assert _norm_rate("0") == "0"
        assert _norm_rate("abc") is None


class TestResolveTaxRate:
    def test_input_consistent(self, tmp_path):
        recs = [{"项目名称": "A", "税率": "0.13"},
                {"项目名称": "B", "税率": "0.13"}]
        assert resolve_tax_rate(recs, interactive=False) == "0.13"

    def test_default_value(self, tmp_path):
        recs = [{"项目名称": "A", "税率": ""}]
        assert resolve_tax_rate(recs, default_value="13", interactive=False) == "0.13"

    def test_saved_default(self, tmp_path):
        mp = os.path.join(tmp_path, "default_tax_rate.json")
        save_default_rate("0.09", mp)
        recs = [{"项目名称": "A", "税率": ""}]
        assert resolve_tax_rate(recs, interactive=False, rate_path=mp) == "0.09"

    def test_interactive_saves(self, tmp_path):
        mp = os.path.join(tmp_path, "default_tax_rate.json")
        recs = [{"项目名称": "A", "税率": ""}]
        from unittest import mock
        # 输入税率 13 → 设为默认 Y
        with mock.patch("builtins.input", side_effect=["13", "Y"]):
            assert resolve_tax_rate(recs, interactive=True, rate_path=mp) == "0.13"
        assert load_default_rate(mp) == "0.13"

    def test_interactive_no_save(self, tmp_path):
        mp = os.path.join(tmp_path, "default_tax_rate.json")
        recs = [{"项目名称": "A", "税率": ""}]
        from unittest import mock
        with mock.patch("builtins.input", side_effect=["0.13", "N"]):
            assert resolve_tax_rate(recs, interactive=True, rate_path=mp) == "0.13"
        assert load_default_rate(mp) is None

    def test_noninteractive_no_rate_raises(self, tmp_path):
        """新用户（无默认税率、无 --tax-rate）非交互运行必须报错，不能默认13%。"""
        mp = os.path.join(tmp_path, "none.json")
        recs = [{"项目名称": "A", "税率": ""}]
        import pytest
        with pytest.raises(ValueError):
            resolve_tax_rate(recs, interactive=False, rate_path=mp)

    def test_six_percent_supported(self, tmp_path):
        """其他税率（如6%）应被支持，不绑定13%。"""
        mp = os.path.join(tmp_path, "default_tax_rate.json")
        save_default_rate("0.06", mp)
        recs = [{"项目名称": "A", "税率": ""}]
        assert resolve_tax_rate(recs, interactive=False, rate_path=mp) == "0.06"


class TestOptionalCols:
    def test_input_consistent(self, tmp_path):
        mp = os.path.join(tmp_path, "optional_cols_prefs.json")
        recs = [{"项目名称": "A", "规格型号": "X1"}, {"项目名称": "B", "规格型号": "X1"}]
        out = resolve_optional_cols(recs, ["规格型号"], interactive=False, prefs_path=mp)
        assert out == {"规格型号": "X1"}

    def test_interactive_skip(self, tmp_path):
        mp = os.path.join(tmp_path, "optional_cols_prefs.json")
        recs = [{"项目名称": "A", "规格型号": ""}]
        from unittest import mock
        with mock.patch("builtins.input", return_value="N"):
            out = resolve_optional_cols(recs, ["规格型号"], interactive=True, prefs_path=mp)
        assert out == {"规格型号": ""}
        assert load_prefs(mp) == {"规格型号": ""}

    def test_interactive_value(self, tmp_path):
        mp = os.path.join(tmp_path, "optional_cols_prefs.json")
        recs = [{"项目名称": "A", "规格型号": ""}]
        from unittest import mock
        with mock.patch("builtins.input", side_effect=["Y", "标准"]):
            out = resolve_optional_cols(recs, ["规格型号"], interactive=True, prefs_path=mp)
        assert out == {"规格型号": "标准"}
        assert load_prefs(mp) == {"规格型号": "标准"}

    def test_pref_respected(self, tmp_path):
        mp = os.path.join(tmp_path, "optional_cols_prefs.json")
        save_prefs({"规格型号": ""}, mp)
        recs = [{"项目名称": "A", "规格型号": ""}]
        # 已有偏好（跳过），不再交互
        out = resolve_optional_cols(recs, ["规格型号"], interactive=True, prefs_path=mp)
        assert out == {"规格型号": ""}


class TestLookupCatalog:
    def test_no_catalog_file(self):
        # 内置编码表已删除，所有查询都返回 None
        assert _lookup_catalog("机械配件") is None
        assert _lookup_catalog("316不锈钢螺母") is None
        assert _lookup_catalog("完全无关的xyz") is None


class TestHeaderAliasMingxi:
    """明细导入表头别名识别（含入库类常见列名）。"""

    def test_common(self):
        assert match_column("品名", config.MINGXI_COLUMNS) == "项目名称"
        assert match_column("单位说明", config.MINGXI_COLUMNS) == "单位"
        assert match_column("入库数量", config.MINGXI_COLUMNS) == "商品数量"
        assert match_column("含税单价", config.MINGXI_COLUMNS) == "商品单价"
        assert match_column("未税单价", config.MINGXI_COLUMNS) == "商品单价"
        assert match_column("入库含税原币金额", config.MINGXI_COLUMNS) == "金额"
        assert match_column("入库税前原币金额", config.MINGXI_COLUMNS) == "金额"

    def test_kaipiao_unaffected(self):
        # 开票信息维护用「单价」仍映射到「单价」，不受明细新增别名影响
        assert match_column("单价") == "单价"
        assert match_column("数量") == "商品数量"

    def test_containment_variants(self):
        """换词也能识别（包含匹配）。"""
        cols = config.MINGXI_COLUMNS
        assert match_column("发货数量", cols) == "商品数量"
        assert match_column("出库数量", cols) == "商品数量"
        assert match_column("个数", cols) == "商品数量"
        assert match_column("件数", cols) == "商品数量"
        assert match_column("售价", cols) == "商品单价"
        assert match_column("总价", cols) == "金额"  # 不能映射到单价
        assert match_column("总额", cols) == "金额"
        assert match_column("货款", cols) == "金额"
        assert match_column("税点", cols) == "税率"

    def test_shangpin_name_with_digits_not_numeric(self, tmp_path):
        """商品名「320偏心环」含数字+文字，不应被误判为数量/单价。"""
        from utils.excel_reader import read_input
        p = os.path.join(tmp_path, "in.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["商品", "商品数量", "商品单价", "金额"])
        ws.append(["320偏心环", 30, 54, 1620])
        ws.append(["255原点盖", 5, 28, 140])
        wb.save(p)
        recs = read_input(p, columns=config.MINGXI_COLUMNS)
        assert len(recs) == 2
        assert recs[0]["项目名称"] == "320偏心环"   # 含数字+文字 → 项目名称
        assert recs[0]["商品数量"] == 30
        assert recs[0]["商品单价"] == 54
        assert recs[0]["金额"] == 1620


class TestAmountDisambiguation:
    """多列命中同一模板列时，用 数量×单价=金额 消歧。"""

    def _make_input(self, path, headers, rows):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        for row in rows:
            ws.append(row)
        wb.save(path)

    def test_total_price_goes_to_amount(self, tmp_path):
        """表头是「品名/数量/单价/总价」时，总价应识别为金额而非单价。"""
        p = os.path.join(tmp_path, "in.xlsx")
        self._make_input(p, ["品名", "数量", "单价", "总价"], [["机床", 3, 100, 300]])
        recs = read_input(p, columns=config.MINGXI_COLUMNS)
        assert len(recs) == 1
        assert recs[0]["商品数量"] == 3
        assert recs[0]["商品单价"] == 100
        assert recs[0]["金额"] == 300

    def test_dual_price_dual_amount_consistent(self, tmp_path):
        """含税/未税双单价+双金额：用 price_choice 指定口径，保留对应一致的一对。"""
        p = os.path.join(tmp_path, "in.xlsx")
        self._make_input(
            p,
            ["品名", "数量", "含税单价", "未税单价", "含税金额", "税前金额"],
            [["轴", 10, 113, 100, 1130, 1000]],
        )
        # 指定不含税 → 单价 100、金额 1000
        recs = read_input(p, columns=config.MINGXI_COLUMNS, price_choice="不含税")
        assert len(recs) == 1
        r = recs[0]
        assert r["商品单价"] == 100
        assert r["金额"] == 1000
        assert abs(100 * 10 - 1000) < 0.01

        # 指定含税 → 单价 113、金额 1130
        recs2 = read_input(p, columns=config.MINGXI_COLUMNS, price_choice="含税")
        assert len(recs2) == 1
        r2 = recs2[0]
        assert r2["商品单价"] == 113
        assert r2["金额"] == 1130
        assert abs(113 * 10 - 1130) < 0.01

    def test_dual_price_interactive_prompts(self, tmp_path):
        """含税/不含税并存且未指定时，会询问用户（模拟输入选择不含税）。"""
        p = os.path.join(tmp_path, "in.xlsx")
        self._make_input(
            p,
            ["品名", "数量", "含税单价", "未税单价", "含税金额", "税前金额"],
            [["轴", 10, 113, 100, 1130, 1000]],
        )
        from unittest import mock
        with mock.patch("builtins.input", side_effect=["2"]):
            recs = read_input(p, columns=config.MINGXI_COLUMNS)
        assert len(recs) == 1
        assert recs[0]["商品单价"] == 100   # 选了不含税
        assert recs[0]["金额"] == 1000


class TestBuildMingxi:
    def test_build_output(self, tmp_path):
        inp = os.path.join(tmp_path, "in.xlsx")
        out_dir = os.path.join(tmp_path, "out")
        _make_input(inp)
        out = build_mingxi(inp, output_dir=out_dir,
                           tax_rate="0.13", code_override="1099900000000000000",
                           interactive=False)
        assert os.path.exists(out)
        wb = openpyxl.load_workbook(out)
        ws = wb[config.MINGXI_SHEET]
        # 表头保留
        assert ws["A3"].value == "项目名称"
        assert ws["H3"].value == "税率"
        # 数据
        assert ws["A4"].value == "320偏心环"
        assert ws["D4"].value == "件"
        assert ws["E4"].value == 30
        assert ws["F4"].value == 54
        assert ws["G4"].value == 1620
        assert ws["H4"].value == "0.13"
        assert ws["B4"].value == "1099900000000000000"
        assert ws["A6"].value is None  # 只有2条
        # 隐藏 sheet 保留
        assert "excelVersion" in wb.sheetnames
        assert "xzqhdm" in wb.sheetnames


class TestValidate:
    def test_missing_required(self):
        recs = [{"项目名称": "A", "金额": "", "税率": ""}]
        errs = validate(recs)
        assert any("金额" in e for e in errs)
        assert any("税率" in e for e in errs)

    def test_ok(self):
        recs = [{"项目名称": "A", "金额": 100, "税率": "0.13"}]
        assert validate(recs) == []