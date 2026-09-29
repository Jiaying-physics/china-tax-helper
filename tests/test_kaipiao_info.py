# -*- coding: utf-8 -*-
"""开票信息维护模块单元测试。"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import config
from utils import excel_reader
from utils.excel_reader import match_column, read_input, normalize_records
from processors.kaipiao_info import build_kaipiao, validate
from processors.tax_code_resolver import resolve_tax_included, resolve_codes, save_tax_map


class TestHeaderMatch:
    def test_basic_alias(self):
        assert match_column("品名") == "项目名称"
        assert match_column("单位") == "单位"
        assert match_column("单价") == "单价"
        assert match_column("含税标志") == "是否含税"

    def test_whitespace_insensitive(self):
        assert match_column(" 商品名称 ") == "项目名称"
        assert match_column(" 分类编码") == "商品和服务税收分类编码"

    def test_unknown(self):
        assert match_column("随便") is None
        assert match_column(None) is None


def _make_input(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["序号", "品名", "单位", "数量", "单价", "金额"])
    ws.append([1, "320偏心环", "件", 30, 54, 1620])
    ws.append([2, "255原点盖", "件", 345, 28, 9660])
    wb.save(path)


class TestReadInput:
    def test_read(self, tmp_path):
        p = os.path.join(tmp_path, "in.xlsx")
        _make_input(p)
        recs = read_input(p)
        assert len(recs) == 2
        assert recs[0]["项目名称"] == "320偏心环"
        assert recs[0]["单位"] == "件"
        assert recs[0]["单价"] == 54

    def test_normalize_fills_columns(self):
        recs = [{"项目名称": "A", "单位": "件"}]
        out = normalize_records(recs)
        assert len(out) == 1
        assert set(out[0].keys()) == set(config.KAIPIAO_COLUMNS)
        assert out[0]["规格型号"] == ""
        assert out[0]["商品和服务税收分类编码"] == ""


class TestValidate:
    def test_missing_code(self):
        recs = [{"项目名称": "A", "商品和服务税收分类编码": ""}]
        errs = validate(recs)
        assert any("编码" in e for e in errs)

    def test_bad_code(self):
        recs = [{"项目名称": "A", "商品和服务税收分类编码": "123"}]
        errs = validate(recs)
        assert any("19" in e for e in errs)

    def test_ok(self):
        recs = [{"项目名称": "A", "商品和服务税收分类编码": "1099900000000000000"}]
        assert validate(recs) == []


class TestBuildKaipiao:
    def test_build_output(self, tmp_path):
        inp = os.path.join(tmp_path, "in.xlsx")
        out_dir = os.path.join(tmp_path, "out")
        _make_input(inp)
        out = build_kaipiao(inp, output_dir=out_dir,
                            default_code="1099900000000000000",
                            tax_included="N")
        assert os.path.exists(out)
        wb = openpyxl.load_workbook(out)
        ws = wb[config.KAIPIAO_SHEET]
        # 表头应保留
        assert ws["A3"].value == "项目名称"
        assert ws["C3"].value == "商品和服务税收分类编码"
        # 第4行起为我们的数据
        assert ws["A4"].value == "320偏心环"
        assert ws["C4"].value == "1099900000000000000"
        assert ws["E4"].value == "件"
        assert ws["F4"].value == 54
        assert ws["A5"].value == "255原点盖"
        # 是否含税统一为 N
        assert ws["G4"].value == "N"
        assert ws["G5"].value == "N"
        # 示例数据第 6 行应为空（只在 2 条后无更多数据）
        assert ws["A6"].value is None


class TestResolveTaxIncluded:
    def test_input_consistent(self):
        recs = [{"项目名称": "A", "是否含税": "Y"},
                {"项目名称": "B", "是否含税": "Y"}]
        assert resolve_tax_included(recs, interactive=False) == "Y"

    def test_default_value_wins(self):
        recs = [{"项目名称": "A", "是否含税": ""}]
        assert resolve_tax_included(recs, default_value="N", interactive=False) == "N"
        assert resolve_tax_included(recs, default_value="Y", interactive=False) == "Y"

    def test_interactive(self):
        recs = [{"项目名称": "A", "是否含税": ""}]
        from unittest import mock
        with mock.patch("builtins.input", return_value="Y"):
            assert resolve_tax_included(recs, interactive=True) == "Y"

    def test_blank_enter_defaults_N(self):
        recs = [{"项目名称": "A", "是否含税": ""}]
        from unittest import mock
        with mock.patch("builtins.input", return_value=""):
            assert resolve_tax_included(recs, interactive=True) == "N"

    def test_inconsistent_input_falls_back(self):
        recs = [{"项目名称": "A", "是否含税": "Y"},
                {"项目名称": "B", "是否含税": "N"}]
        # 输入列值不一致 → 回退到 default
        assert resolve_tax_included(recs, default_value="N", interactive=False) == "N"


class TestResolveCodes:
    def _tmp_map(self, tmp_path):
        return os.path.join(tmp_path, "tax_code_map.json")

    def test_input_has_code(self, tmp_path):
        mp = self._tmp_map(tmp_path)
        recs = [{"项目名称": "A", "商品和服务税收分类编码": "1099900000000000000"}]
        # 已有编码直接采用，不读记忆库、不交互
        out = resolve_codes(recs, map_path=mp, default_code=None, interactive=True)
        assert out == {"A": "1099900000000000000"}

    def test_map_hit_no_interactive(self, tmp_path):
        mp = self._tmp_map(tmp_path)
        save_tax_map({"螺母": "1099900000000000000"}, mp)
        recs = [{"项目名称": "320锁紧螺母", "商品和服务税收分类编码": ""}]
        out = resolve_codes(recs, map_path=mp, default_code=None, interactive=True)
        assert out == {"320锁紧螺母": "1099900000000000000"}

    def test_default_applies_to_all(self, tmp_path):
        mp = self._tmp_map(tmp_path)
        recs = [{"项目名称": "甲", "商品和服务税收分类编码": ""},
                {"项目名称": "乙", "商品和服务税收分类编码": ""}]
        out = resolve_codes(recs, map_path=mp, default_code="2010500000000000000",
                            interactive=False)
        assert out == {"甲": "2010500000000000000", "乙": "2010500000000000000"}

    def test_batch_interactive_single_code(self, tmp_path):
        mp = self._tmp_map(tmp_path)
        recs = [{"项目名称": "甲", "商品和服务税收分类编码": ""},
                {"项目名称": "乙", "商品和服务税收分类编码": ""}]
        from unittest import mock
        with mock.patch("builtins.input", return_value="1099900000000000000"):
            out = resolve_codes(recs, map_path=mp, interactive=True)
        assert out == {"甲": "1099900000000000000", "乙": "1099900000000000000"}
        # 记忆库应保存
        assert os.path.exists(mp)