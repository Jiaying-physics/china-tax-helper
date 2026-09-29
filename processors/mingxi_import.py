# -*- coding: utf-8 -*-
"""明细导入模块。

以 templates/明细导入.xlsx 为标准模板（只读基准，永不修改），
复制作为输出基底，清空模板自带示例数据，写入输入数据，保持表头/样式一致。
所有数据处理均为本地完成。

交互流程：
  1. 解析税率（必填）：输入自带 / 命令行 / 本地默认 / 询问+记忆
  2. 非必填列（规格型号/折扣金额/优惠政策类型/煤炭种类）：首次询问是否要填，
     要填则问统一值并记忆；输入自带直接用
  3. 商品和服务税收分类编码：输入自带 → 本地记忆库/内置编码表 → 查不到按非必填询问
"""
import json
import os
from typing import Optional

import openpyxl

import config
from .tax_code_resolver import load_tax_map, _lookup
from .tax_rate_resolver import resolve_tax_rate
from .optional_col_resolver import resolve_optional_cols


def _copy_template() -> openpyxl.Workbook:
    """读入模板文件到内存工作簿（保持全部样式），不修改原文件。"""
    src = os.path.join(config.TEMPLATES_DIR, config.MINGXI_TEMPLATE_FILE)
    if not os.path.exists(src):
        raise FileNotFoundError(f"模板文件不存在: {src}")
    return openpyxl.load_workbook(src)


def _clear_template_data(ws) -> None:
    """清空模板自带示例数据（第4行起），仅清内容，保留样式。"""
    for r in range(config.MINGXI_DATA_START, ws.max_row + 1):
        for c in range(1, len(config.MINGXI_COLUMNS) + 1):
            ws.cell(row=r, column=c).value = None


def _write_formatted(ws, row: int, rec: dict) -> None:
    """将一条记录写入第 row 行，值写入对应模板列。"""
    for idx, col in enumerate(config.MINGXI_COLUMNS, start=1):
        v = rec.get(col, "")
        cell = ws.cell(row=row, column=idx)
        cell.value = v if v != "" else None


def _lookup_catalog(name: str) -> Optional[str]:
    """在内置编码表 data/tax_code_catalog.json 中按关键词匹配。返回编码或 None。"""
    catalog_path = os.path.join(config.DATA_DIR, "tax_code_catalog.json")
    if not os.path.exists(catalog_path):
        return None
    with open(catalog_path, encoding="utf-8") as f:
        catalog = json.load(f)
    for code, keywords in catalog.items():
        for kw in keywords:
            if kw and kw in name:
                return code
    return None


def _resolve_codes_local(records, code_override=None, interactive=True) -> dict:
    """商品和服务税收分类编码解析：输入自带 → 记忆库 → 内置编码表 → 询问。"""
    mapping = load_tax_map()  # 复用开票信息维护积累的记忆库
    result = {}
    unresolved = set()

    for rec in records:
        name = str(rec.get("项目名称", "")).strip() or "未知项目"
        raw = rec.get("商品和服务税收分类编码")
        if raw not in (None, ""):
            code = str(raw).strip()
            if code.isdigit() and len(code) == 19:
                result[name] = code
                continue
        found = _lookup(name, mapping)
        if found:
            result[name] = found
            continue
        # 内置编码表兜底
        cat_code = _lookup_catalog(name)
        if cat_code:
            result[name] = cat_code
            continue
        unresolved.add(name)

    if code_override:
        for name in unresolved:
            result[name] = code_override
        return result

    if interactive and unresolved:
        print("\n以下项目未能匹配到本地税收分类编码，请提供（输一个应用到全部，回车跳过）：")
        for idx, name in enumerate(sorted(unresolved), 1):
            print(f"  {idx}. {name}")
        while True:
            ans = input("请输入 19 位税收分类编码应用到以上全部项目（回车=跳过）: ").strip()
            if not ans:
                print("  未提供，编码留空。")
                for name in unresolved:
                    result[name] = ""
                break
            if ans.isdigit() and len(ans) == 19:
                for name in unresolved:
                    result[name] = ans
                    mapping[name] = ans  # 记忆
                break
            print("  编码格式错误，需为 19 位纯数字。")
        from .tax_code_resolver import save_tax_map
        save_tax_map(mapping)
    return result


def validate(records) -> list:
    """校验必填列：项目名称、金额、税率。返回错误列表。"""
    errors = []
    for i, rec in enumerate(records, start=config.MINGXI_DATA_START):
        name = rec.get("项目名称") or ""
        if not name:
            errors.append(f"第{i}行：缺少项目名称")
        amt = rec.get("金额")
        if amt in (None, ""):
            errors.append(f"第{i}行（{name or '无名称'}）：缺少金额")
        rate = rec.get("税率")
        if rate in (None, ""):
            errors.append(f"第{i}行（{name or '无名称'}）：缺少税率")
    return errors


def build_mingxi(
    input_path: str,
    output_dir: Optional[str] = None,
    tax_rate: Optional[str] = None,
    code_override: Optional[str] = None,
    sheet_name: Optional[str] = None,
    price_choice: Optional[str] = None,
    interactive: bool = True,
) -> str:
    """处理输入文件，生成明细导入标准 Excel，返回输出路径。"""
    from utils.excel_reader import read_input, normalize_records

    records = read_input(input_path, columns=config.MINGXI_COLUMNS, sheet_name=sheet_name,
                         price_choice=price_choice)
    records = normalize_records(records, columns=config.MINGXI_COLUMNS)
    if not records:
        raise ValueError("输入文件中没有识别到数据")

    # 1. 税率（必填）
    rate = resolve_tax_rate(records, default_value=tax_rate, interactive=interactive)
    for rec in records:
        rec["税率"] = rate
    print(f"本张发票税率：{rate}")

    # 2. 非必填列统一偏好
    optional_vals = resolve_optional_cols(
        records, config.MINGXI_OPTIONAL_COLS, interactive=interactive
    )
    for col, val in optional_vals.items():
        for rec in records:
            if not rec.get(col):  # 仅当输入未提供该列时填统一值
                rec[col] = val

    # 3. 商品和服务税收分类编码（本地优先，查不到再问）
    codes = _resolve_codes_local(records, code_override=code_override, interactive=interactive)
    for rec in records:
        name = rec.get("项目名称")
        rec["商品和服务税收分类编码"] = codes.get(name, rec.get("商品和服务税收分类编码") or "")

    # 校验
    errors = validate(records)
    if errors:
        print("\n⚠ 校验提示（不阻断输出）：")
        for e in errors:
            print("  -", e)

    wb = _copy_template()
    ws = wb[config.MINGXI_SHEET]
    _clear_template_data(ws)

    for i, rec in enumerate(records):
        _write_formatted(ws, config.MINGXI_DATA_START + i, rec)

    output_dir = output_dir or config.OUTPUTS_DIR
    os.makedirs(output_dir, exist_ok=True)
    out_path = os.path.join(output_dir, config.output_file_name(input_path, config.MINGXI_OUTPUT_PREFIX))
    wb.save(out_path)
    print(f"\n✅ 已生成明细导入表：{out_path}")
    return out_path