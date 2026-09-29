# -*- coding: utf-8 -*-
"""开票信息维护模块。

以 templates/开票信息维护.xlsx 为标准模板（只读基准，永不修改），
复制作为输出基底，清空模板自带示例数据，写入输入数据，保持表头/样式一致。
所有数据处理均为本地完成。
"""
import os
from typing import Optional

import openpyxl

import config
from .tax_code_resolver import resolve_codes, resolve_tax_included


def _copy_template() -> openpyxl.Workbook:
    """读入模板文件到内存工作簿（保持全部样式），不修改原文件。"""
    src = os.path.join(config.TEMPLATES_DIR, config.KAIPIAO_TEMPLATE_FILE)
    if not os.path.exists(src):
        raise FileNotFoundError(f"模板文件不存在: {src}")
    return openpyxl.load_workbook(src)


def _clear_template_data(ws) -> None:
    """清空模板自带示例数据（第4行起），仅清内容，保留样式。"""
    # 模板示例数据到 114 行；为稳妥起见清到模板 max_row
    for r in range(config.KAIPIAO_DATA_START, ws.max_row + 1):
        for c in range(1, len(config.KAIPIAO_COLUMNS) + 1):
            ws.cell(row=r, column=c).value = None


def _write_formatted(ws, row: int, rec: dict) -> None:
    """将一条记录写入第 row 行，值写入对应模板列。样式继承（数据行无值时会显现有边框，由模板自带）。"""
    for idx, col in enumerate(config.KAIPIAO_COLUMNS, start=1):
        v = rec.get(col, "")
        cell = ws.cell(row=row, column=idx)
        cell.value = v if v != "" else None


def validate(records) -> list:
    """校验必填列与编码格式；返回错误信息列表。"""
    errors = []
    for i, rec in enumerate(records, start=config.KAIPIAO_DATA_START):
        name = rec.get("项目名称") or ""
        if not name:
            errors.append(f"第{i}行：缺少项目名称")
        code = rec.get("商品和服务税收分类编码")
        if not code:
            errors.append(f"第{i}行（{name or '无名称'}）：缺少税收分类编码")
        else:
            s = str(code).strip()
            if not (s.isdigit() and len(s) == 19):
                errors.append(f"第{i}行（{name}）：税收分类编码需为19位数字，当前为 {code!r}")
    return errors


def build_kaipiao(
    input_path: str,
    output_dir: Optional[str] = None,
    default_code: Optional[str] = None,
    tax_included: Optional[str] = None,
    sheet_name: Optional[str] = None,
    price_choice: Optional[str] = None,
    interactive: bool = True,
) -> str:
    """处理输入文件，生成开票信息维护标准 Excel，返回输出路径。"""
    from utils.excel_reader import read_input, normalize_records

    records = read_input(input_path, sheet_name=sheet_name, price_choice=price_choice)
    records = normalize_records(records)
    if not records:
        raise ValueError("输入文件中没有识别到数据")

    # 先解析整表统一的「是否含税」标识（先问含税）
    included = resolve_tax_included(records, default_value=tax_included, interactive=interactive)
    for rec in records:
        rec["是否含税"] = included
    print(f"本张发票是否含税标识：{'Y（含税）' if included == 'Y' else 'N（不含税）'}")

    # 解析税收分类编码（后问编码）
    codes = resolve_codes(records, default_code=default_code, interactive=interactive)

    # 将解析到的编码写回记录
    for rec in records:
        name = rec.get("项目名称")
        rec["商品和服务税收分类编码"] = codes.get(name, rec.get("商品和服务税收分类编码") or "")

    # 校验
    errors = validate(records)
    if errors:
        print("\n⚠ 校验提示（不阻断输出）：")
        for e in errors:
            print("  -", e)

    wb = _copy_template()
    ws = wb[config.KAIPIAO_SHEET]
    _clear_template_data(ws)

    for i, rec in enumerate(records):
        _write_formatted(ws, config.KAIPIAO_DATA_START + i, rec)

    output_dir = output_dir or config.OUTPUTS_DIR
    os.makedirs(output_dir, exist_ok=True)
    out_path = os.path.join(output_dir, config.output_file_name(input_path))
    wb.save(out_path)
    print(f"\n✅ 已生成开票信息维护表：{out_path}")
    return out_path